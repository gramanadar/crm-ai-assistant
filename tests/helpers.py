from io import BytesIO

import pymupdf
from docx import Document

from tests.conftest import APPROVER


def docx_bytes(paragraphs: list[tuple[str, str]], tables: list[list[list[str]]] | None = None) -> bytes:
    document = Document()
    for style, text in paragraphs:
        document.add_paragraph(text, style=style)
    for rows in tables or []:
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for row_index, row in enumerate(rows):
            for col_index, value in enumerate(row):
                table.rows[row_index].cells[col_index].text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pdf_bytes(pages: list[list[tuple[str, int]]]) -> bytes:
    document = pymupdf.open()
    for lines in pages:
        page = document.new_page()
        y = 72
        for text, size in lines:
            page.insert_text((72, y), text, fontsize=size)
            y += size + 12
    buffer = BytesIO()
    document.save(buffer)
    document.close()
    return buffer.getvalue()


def upload(client, filename: str, data: bytes, headers=None):
    return client.post(
        "/v1/documents",
        files={"file": (filename, data, "application/octet-stream")},
        headers=headers or APPROVER,
    )
